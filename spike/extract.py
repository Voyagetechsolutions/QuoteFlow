"""
Extraction spike: messy freight rate sheet (.xlsx / .pdf) -> canonical JSON rows.

This is THROWAWAY code whose only job is to answer one question before any app
is built: can we get usable structured rate rows out of a real, messy sheet,
with the bad cases landing as fast 30-second fixes rather than dead ends?

Design choices that mirror the PRD:
  - Every emitted row carries `needs_review` + `issues`, because extraction is
    never 100% and the editable review table is the product, not a fallback.
  - Free-text validity ("Rate Validity: 01 Jun ...") and section headers
    ("SEA FREIGHT - FCL") are context that gets carried DOWN onto each row.
  - Excel path: openpyxl. PDF path: pdfplumber table extraction.

Usage:
    python spike/extract.py samples/sample_rate_sheet.xlsx
    python spike/extract.py samples/sample_rate_sheet.pdf
    python spike/extract.py <file> --pretty
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field, asdict
from datetime import date
from typing import Optional

import openpyxl
import pdfplumber

from normalize import normalize_basis

# --- canonical schema --------------------------------------------------------

CHARGE_TYPE_BY_SECTION = {
    "SEA FREIGHT": "Ocean Freight",
    "OCEAN FREIGHT": "Ocean Freight",
    "ROAD FREIGHT": "Road Freight",
    "AIR FREIGHT": "Air Freight",
    "LOCAL CHARGES": "Local Charge",
    "LOCAL CHARGE": "Local Charge",
}

HEADER_TOKENS = {"origin", "destination", "dest", "unit", "rate", "currency",
                 "ccy", "remarks", "remark", "lane", "charge", "amount",
                 "pol", "pod", "basis", "transit", "description", "price",
                 "from", "to", "service", "tariff"}


def container_basis(header: Optional[str]) -> Optional[str]:
    """A container-size column header (20'GP, 40HC, 20ft, ...) -> basis enum.

    Real FCL rate sheets put the rate in per-container-size columns rather than
    a single 'rate' column, so each such column is its own rate with a known
    basis. Returns None for non-container headers.
    """
    h = re.sub(r"[^a-z0-9]", "", (header or "").lower())
    if re.fullmatch(r"20(gp|dc|hc|hq|ft|rf|reefer)?", h):
        return "per_container_20"
    if re.fullmatch(r"40(gp|dc|hc|hq|ft|rf|reefer)?", h):
        return "per_container_40"
    return None

CURRENCY_SYMBOL = {"$": "USD", "R": "ZAR", "€": "EUR", "£": "GBP"}
KNOWN_CURRENCIES = {"USD", "ZAR", "EUR", "GBP", "ZWL", "BWP", "ZMW", "MWK"}

# free-text validity line, e.g. "Rate Validity: 01 Jun 2026 - 31 Aug 2026"
_DATE = r"(\d{1,2}\s+\w{3,9}\s+\d{4})"
VALIDITY_RE = re.compile(_DATE + r"\s*[-–to]+\s*" + _DATE, re.IGNORECASE)

# rate value, e.g. "USD 1,850.00", "$3,400", "4 050", "ZAR 28,500"
RATE_RE = re.compile(
    r"(?P<cur>[A-Z]{3}|[$R€£])?\s*"
    r"(?P<amt>\d[\d\s,]*(?:\.\d+)?)\s*"
    r"(?P<cur2>[A-Z]{3})?",
)


@dataclass
class RateRow:
    charge_type: Optional[str] = None
    lane_origin: Optional[str] = None
    lane_destination: Optional[str] = None
    unit: Optional[str] = None
    rate: Optional[float] = None
    currency: Optional[str] = None
    valid_from: Optional[str] = None
    valid_to: Optional[str] = None
    remark: Optional[str] = None
    # filled by the normalization post-pass (spike/normalize.py)
    charge_code: Optional[str] = None
    charge_label: Optional[str] = None
    basis: Optional[str] = None
    confidence: float = 1.0
    needs_review: bool = False
    issues: list[str] = field(default_factory=list)
    source: str = ""  # raw cell text, for the review table

    def flag(self, msg: str) -> None:
        self.needs_review = True
        if msg not in self.issues:
            self.issues.append(msg)


# --- shared parsing helpers --------------------------------------------------

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun",
     "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


def parse_loose_date(text: str) -> Optional[str]:
    """'01 Jun 2026' -> '2026-06-01' (ISO). Returns None if unparseable."""
    m = re.match(r"(\d{1,2})\s+(\w{3,9})\s+(\d{4})", text.strip())
    if not m:
        return None
    day, mon, year = m.groups()
    month = _MONTHS.get(mon[:3].lower())
    if not month:
        return None
    try:
        return date(int(year), month, int(day)).isoformat()
    except ValueError:
        return None


def parse_validity(text: str) -> tuple[Optional[str], Optional[str]]:
    m = VALIDITY_RE.search(text or "")
    if not m:
        return None, None
    return parse_loose_date(m.group(1)), parse_loose_date(m.group(2))


def parse_rate(text: str) -> tuple[Optional[float], Optional[str]]:
    """'USD 1,850.00' -> (1850.0, 'USD'); '$3,400' -> (3400.0, 'USD')."""
    if text is None:
        return None, None
    s = str(text).strip()
    if not s:
        return None, None
    # Non-numeric placeholders — don't fabricate a rate from "see note 2",
    # "on application", "POA", "TBD", etc. A flagged blank beats a wrong number.
    if re.search(r"(?i)\b(see|note|application|request|poa|tba|tbd|n/?a|nil)\b", s):
        return None, None
    m = RATE_RE.search(s)
    if not m:
        return None, None
    raw_amt = m.group("amt").replace(",", "").replace(" ", "")
    try:
        amount = float(raw_amt)
    except ValueError:
        return None, None
    cur = m.group("cur") or m.group("cur2")
    if cur in CURRENCY_SYMBOL:
        cur = CURRENCY_SYMBOL[cur]
    if cur and cur not in KNOWN_CURRENCIES:
        cur = None
    return amount, cur


def charge_type_for(section: Optional[str]) -> Optional[str]:
    if not section:
        return None
    up = section.upper()
    for key, val in CHARGE_TYPE_BY_SECTION.items():
        if key in up:
            return val
    return section.title()


def is_header_row(cells: list[str]) -> bool:
    toks = [c.strip().lower() for c in cells if c and c.strip()]
    if not toks:
        return False
    hits = sum(1 for t in toks if t in HEADER_TOKENS)
    hits += sum(1 for c in cells if container_basis(c))
    return hits >= 2


def looks_like_footnote(cells: list[str]) -> bool:
    joined = " ".join(c for c in cells if c).strip().lower()
    return bool(joined) and joined.startswith(
        ("rates exclude", "subject to", "usd rates", "note", "terms", "e&oe"))


def is_section_row(cells: list[str]) -> Optional[str]:
    """A lone non-empty cell in ALL CAPS / known section keyword."""
    nonempty = [c.strip() for c in cells if c and c.strip()]
    if len(nonempty) != 1:
        return None
    text = nonempty[0]
    up = text.upper()
    if any(k in up for k in CHARGE_TYPE_BY_SECTION) or up == text:
        return text
    return None


# --- column mapping ----------------------------------------------------------

def build_colmap(header_cells: list[str]) -> dict:
    """Classify header columns. Returns field indices plus `rate_cols`: a list
    of (index, basis) where basis is a container enum for size columns or None
    for a generic 'rate'/'amount' column."""
    cm: dict = {"rate_cols": []}
    for idx, cell in enumerate(header_cells):
        t = (cell or "").strip().lower()
        cb = container_basis(cell)
        if cb:
            cm["rate_cols"].append((idx, cb))
        elif t in ("origin", "from", "pol", "port of loading", "loading port"):
            cm["origin"] = idx
        elif t in ("destination", "dest", "to", "pod",
                   "port of discharge", "discharge port"):
            cm["destination"] = idx
        elif t in ("charge", "charge type", "description", "item", "service"):
            cm["charge"] = idx
        elif t in ("basis", "unit", "per"):
            cm["basis"] = idx
        elif t in ("rate", "amount", "price", "tariff"):
            cm["rate_cols"].append((idx, None))
        elif t in ("currency", "ccy", "cur"):
            cm["currency"] = idx
        elif t in ("remarks", "remark", "notes", "note", "transit"):
            cm["remark"] = idx
    return cm


def rows_from_cells(cells: list[str], colmap: dict, ctx: dict) -> list[RateRow]:
    """Build RateRow(s) from one data row. Wide container tables (rate split
    across 20'/40' columns) yield ONE row per size column (an unpivot)."""
    def get(key: str) -> Optional[str]:
        i = colmap.get(key)
        if i is None or i >= len(cells):
            return None
        v = cells[i]
        return v.strip() if isinstance(v, str) else (str(v) if v is not None else None)

    origin = get("origin")
    charge = get("charge")
    basis_text = get("basis")
    rate_cols = colmap.get("rate_cols", [])
    # A data row needs a lane/charge identity and at least one rate column.
    if not origin and not charge:
        return []

    # An explicit non-container basis (per B/L, per shipment, %) applies to the
    # whole row; a generic "per container" defers to the size column.
    row_basis = normalize_basis(basis_text) if basis_text else None
    if row_basis in ("per_container_20", "per_container_40"):
        row_basis = None

    source = " | ".join(c for c in cells if c and str(c).strip())
    out: list[RateRow] = []
    for idx, col_basis in rate_cols:
        rate_text = get_at(cells, idx)
        if not rate_text:
            continue
        row = RateRow(
            charge_type=charge or ctx.get("section_charge_type"),
            lane_origin=origin or None,
            lane_destination=get("destination") or None,
            unit=basis_text or None,
            remark=get("remark") or None,
            basis=row_basis or col_basis,
            valid_from=ctx.get("valid_from"),
            valid_to=ctx.get("valid_to"),
            source=source,
        )
        amount, cur = parse_rate(rate_text)
        row.rate = amount
        col_cur = get("currency")
        row.currency = (
            col_cur.upper() if col_cur and col_cur.upper() in KNOWN_CURRENCIES
            else cur
        )
        validate_row(row, raw_rate_text=rate_text)
        out.append(row)
    return out


def get_at(cells: list[str], i: int) -> Optional[str]:
    if i is None or i >= len(cells):
        return None
    v = cells[i]
    return v.strip() if isinstance(v, str) else (str(v) if v is not None else None)


def validate_row(row: RateRow, raw_rate_text: Optional[str] = None) -> RateRow:
    """Shared review-flagging. Both the table path and the vision path call this
    so a flag means the same thing regardless of which extractor produced it."""
    if row.rate is None:
        row.flag(f"could not parse rate from {raw_rate_text!r}"
                 if raw_rate_text else "rate missing or unparseable")
    if row.currency is None:
        row.flag("currency unknown")
    if not row.lane_origin:
        row.flag("missing origin")
    if row.charge_type != "Local Charge" and not row.lane_destination:
        row.flag("missing destination")
    if not row.valid_from or not row.valid_to:
        row.flag("validity not found on sheet")
    return row


# --- Excel path --------------------------------------------------------------

def _extract_from_grid(grid: list[list[str]]) -> list[RateRow]:
    """Core row-by-row parse, shared by the Excel and CSV paths."""
    rows: list[RateRow] = []
    ctx: dict = {}
    colmap: dict[str, int] = {}

    for cells in grid:
        vf, vt = parse_validity(" ".join(cells))
        if vf or vt:
            ctx["valid_from"], ctx["valid_to"] = vf, vt
            continue
        if looks_like_footnote(cells):
            continue
        section = is_section_row(cells)
        if section:
            ctx["section_charge_type"] = charge_type_for(section)
            colmap = {}  # header repeats per section
            continue
        if is_header_row(cells):
            colmap = build_colmap(cells)
            continue
        if not colmap:
            continue
        rows.extend(rows_from_cells(cells, colmap, ctx))
    return rows


def extract_xlsx(path: str) -> list[RateRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    grid = [
        ["" if c is None else str(c) for c in raw]
        for raw in ws.iter_rows(values_only=True)
    ]
    return _extract_from_grid(grid)


def extract_csv(path: str) -> list[RateRow]:
    import csv

    with open(path, newline="", encoding="utf-8-sig") as fh:
        grid = [[(c or "").strip() for c in row] for row in csv.reader(fh)]
    return _extract_from_grid(grid)


# --- PDF path ----------------------------------------------------------------

def extract_pdf(path: str) -> list[RateRow]:
    rows: list[RateRow] = []
    ctx: dict = {}

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            # carry free-text validity from the page text
            text = page.extract_text() or ""
            vf, vt = parse_validity(text)
            if vf or vt:
                ctx.setdefault("valid_from", vf)
                ctx.setdefault("valid_to", vt)

            # Section headings are separate objects floating above their tables.
            # Record each heading's vertical position so we can bind a table to
            # the nearest heading ABOVE it (the part Excel gives us for free).
            sections = _sections_with_positions(page)

            for table in page.find_tables():
                data = table.extract()
                if not data:
                    continue
                header = ["" if c is None else str(c) for c in data[0]]
                if not is_header_row(header):
                    continue
                colmap = build_colmap(header)

                table_top = table.bbox[1]  # (x0, top, x1, bottom)
                section = _section_above(sections, table_top)
                ctx["section_charge_type"] = (
                    charge_type_for(section) if section else None)

                for raw in data[1:]:
                    cells = ["" if c is None else str(c) for c in raw]
                    if is_header_row(cells) or looks_like_footnote(cells):
                        continue
                    rows.extend(rows_from_cells(cells, colmap, ctx))
    return rows


def _sections_with_positions(page) -> list[tuple[float, str]]:
    """Reconstruct text lines with their y-top; keep those that name a section.

    Returns [(top, section_text), ...] sorted top-to-bottom.
    """
    words = page.extract_words(use_text_flow=False)
    lines: dict[float, list] = {}
    for w in words:
        # bucket words by rounded vertical position into lines
        key = round(w["top"] / 2) * 2
        lines.setdefault(key, []).append(w)

    found: list[tuple[float, str]] = []
    for top, ws in lines.items():
        ws.sort(key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in ws).strip()
        up = text.upper()
        if any(k in up for k in CHARGE_TYPE_BY_SECTION):
            found.append((float(top), text))
    found.sort()
    return found


def _section_above(sections: list[tuple[float, str]],
                   table_top: float) -> Optional[str]:
    """Nearest section heading whose y-position is above the table top."""
    candidate = None
    for top, text in sections:
        if top <= table_top:
            candidate = text
        else:
            break
    return candidate


# --- entrypoint --------------------------------------------------------------

def extract(path: str) -> list[RateRow]:
    lower = path.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xls")):
        return extract_xlsx(path)
    if lower.endswith(".csv"):
        return extract_csv(path)
    if lower.endswith(".pdf"):
        return extract_pdf(path)
    raise SystemExit(f"unsupported file type: {path}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Rate sheet -> canonical JSON rows")
    ap.add_argument("path", help="path to .xlsx or .pdf rate sheet")
    ap.add_argument("--pretty", action="store_true", help="indent JSON output")
    args = ap.parse_args()

    rows = extract(args.path)
    payload = {
        "source_file": args.path,
        "row_count": len(rows),
        "needs_review_count": sum(1 for r in rows if r.needs_review),
        "rows": [asdict(r) for r in rows],
    }
    json.dump(payload, sys.stdout,
              indent=2 if args.pretty else None, ensure_ascii=False)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
